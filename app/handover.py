"""Rendering the hand-over .docx from templates_docx/handover_template.docx
and logging each generated form to handover_records (the hand-over history)."""

from __future__ import annotations

import datetime as dt
import io
import json
import os
import re
from copy import deepcopy

import segno
from docx.oxml import parse_xml
from docx.shared import Mm, Pt
from docxtpl import DocxTemplate

from .db import get_connection, log_activity
from .paths import get_handover_template_path, get_handovers_dir

# ASSIGNMENT/TEMP hand equipment OUT to the employee; RETURN gets it back IN.
# The (Hand Over)/(Receive) suffixes make that direction explicit in the UI
# without needing a separate data field.
HO_TYPES = [
    ("ASSIGNMENT", "ASSIGNMENT (Hand Over)"),
    ("TEMP", "TEMP ASSIGNMENT (Hand Over)"),
    ("RETURN", "RETURN (Receive)"),
    ("OTHER", "OTHER"),
]
REASONS = ["NEWCOMER", "ROTATE", "MATERNITY", "RESIGNED", "OTHER"]
ASSET_CONDITIONS = ["NEW", "WAREHOUSE"]

CHECKED = "☑"
UNCHECKED = "☐"


def _box(selected: bool) -> str:
    return CHECKED if selected else UNCHECKED


# Every generated form carries a QR code (page header, top-left - the
# Shinhan logo already owns top-right) so a scanned, signed copy can be
# matched back to its handover_records row automatically. Payload, all
# ASCII, "|"-separated, version first so the format can change later:
#   AMHO1|<record id>|<ho_date YYYY-MM-DD>|<user_no>|<ho_type>
# The id alone identifies the form; date/user/type let a reader sort the
# scan (and sanity-check the id) without access to this database.
QR_PREFIX = "AMHO1"
QR_SIZE = Mm(18)
QR_TOP = Mm(6)  # from the page edge - clear of most printers' dead zone


def form_number(record_id: int) -> str:
    """Human-readable form number printed beside the QR and shown in the
    History page - the same id the QR carries."""
    return f"HO-{int(record_id):06d}"


def _ascii_field(value) -> str:
    return re.sub(r"[^0-9A-Za-z._-]", "", str(value or ""))


def qr_payload(record_id: int, ho_date: dt.date, user_no, ho_type) -> str:
    return "|".join(
        [QR_PREFIX, str(int(record_id)), ho_date.isoformat(), _ascii_field(user_no),
         _ascii_field((ho_type or "").upper())]
    )


def _add_header_qr(doc, payload: str, number: str) -> None:
    """Puts the QR (floating, anchored to the page's top-left margin, text
    wrapping around it like the logo opposite) and the form number into the
    header shared by every page, so each page of a multi-page form - or of a
    batch scan - identifies itself. Added to the in-memory copy docxtpl
    renders from; the template file itself is unchanged."""
    buf = io.BytesIO()
    segno.make_qr(payload, error="m").save(buf, kind="png", scale=10, border=2)
    buf.seek(0)

    paragraph = doc.sections[0].header.paragraphs[0]
    pic_run = paragraph.add_run()
    pic_run.add_picture(buf, width=QR_SIZE, height=QR_SIZE)
    inline = pic_run._r.find(".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline")
    anchor = parse_xml(
        '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'distT="0" distB="0" distL="0" distR="72000" simplePos="0" relativeHeight="251660288" '
        'behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="margin"><wp:align>left</wp:align></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{int(QR_TOP)}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{int(QR_SIZE)}" cy="{int(QR_SIZE)}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        '<wp:wrapSquare wrapText="bothSides"/>'
        "</wp:anchor>"
    )
    # Reuse the inline's docPr/graphic (they carry the image relationship id).
    for tag in ("docPr", "cNvGraphicFramePr"):
        el = inline.find(f"{{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}}{tag}")
        if el is not None:
            anchor.append(deepcopy(el))
    anchor.append(deepcopy(inline.find("{http://schemas.openxmlformats.org/drawingml/2006/main}graphic")))
    anchor.find("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr").set(
        "name", "HandoverQR"
    )
    inline.getparent().replace(inline, anchor)

    num_run = paragraph.add_run(number)
    num_run.font.name = "Times New Roman"
    num_run.font.size = Pt(8)


def resolve_ho_date(ho_date: str | None) -> dt.date:
    """`ho_date` is an HTML <input type=date> value ("YYYY-MM-DD"). Falls
    back to today (the system date) if missing or unparseable."""
    if ho_date:
        try:
            return dt.datetime.strptime(ho_date, "%Y-%m-%d").date()
        except ValueError:
            pass
    return dt.date.today()


def render_handover_docx(data: dict, record_id: int) -> str:
    """`data` keys: ict_rep_name, ict_rep_id, receiving_name, receiving_title,
    receiving_dept, receiving_id, ho_type, temp_due_date, other_type_text,
    reason, other_reason_text, user_no, ho_date ("YYYY-MM-DD", defaults to
    today), signature_receiving_name, signature_prepared_by, assets (list of
    dict with equipment_name, qty, serial_number, status). `record_id` is
    the form's handover_records id, encoded in the header QR. Returns the
    saved .docx path.
    """
    now = dt.datetime.now()
    ho_date = resolve_ho_date(data.get("ho_date"))
    assets = []
    for i, a in enumerate(data.get("assets", []), start=1):
        assets.append(
            {
                "no": i,
                "equipment_name": a.get("equipment_name", ""),
                "qty": a.get("qty", 1),
                "serial_number": a.get("serial_number", ""),
                "status": a.get("status", ""),
            }
        )
    total_qty = sum(int(a["qty"] or 1) for a in assets)

    ho_type = (data.get("ho_type") or "").upper()
    reason = (data.get("reason") or "").upper()

    ctx = {
        "ho_day": ho_date.strftime("%d"),
        "ho_month": ho_date.strftime("%m"),
        "ho_year": ho_date.strftime("%Y"),
        "ict_rep_name": data.get("ict_rep_name", ""),
        "ict_rep_id": data.get("ict_rep_id", ""),
        "receiving_name": data.get("receiving_name", ""),
        "receiving_title": data.get("receiving_title", ""),
        "receiving_dept": data.get("receiving_dept", ""),
        "receiving_id": data.get("receiving_id", ""),
        "type_assignment_box": _box(ho_type == "ASSIGNMENT"),
        "type_temp_box": _box(ho_type == "TEMP"),
        "type_return_box": _box(ho_type == "RETURN"),
        "type_other_box": _box(ho_type == "OTHER"),
        "temp_due_date": data.get("temp_due_date", ""),
        "other_type_text": data.get("other_type_text", ""),
        "reason_newcomer_box": _box(reason == "NEWCOMER"),
        "reason_rotate_box": _box(reason == "ROTATE"),
        "reason_maternity_box": _box(reason == "MATERNITY"),
        "reason_resigned_box": _box(reason == "RESIGNED"),
        "reason_other_box": _box(reason == "OTHER"),
        "other_reason_text": data.get("other_reason_text", ""),
        "assets": assets,
        "total_qty": total_qty,
        "signature_receiving_name": data.get("signature_receiving_name") or data.get("receiving_name", ""),
        "signature_prepared_by": data.get("signature_prepared_by") or data.get("ict_rep_name", ""),
    }

    tpl = DocxTemplate(get_handover_template_path())
    number = form_number(record_id)
    # Before render, not after: docxtpl re-reads each header's XML and swaps
    # in a fresh part, so edits made to the header afterwards are lost.
    _add_header_qr(
        tpl.get_docx(), qr_payload(record_id, ho_date, data.get("user_no"), ho_type), number
    )
    tpl.render(ctx)

    safe_user = "".join(c for c in str(data.get("user_no") or "unknown") if c.isalnum()) or "unknown"
    filename = f"HANDOVER_{number}_{safe_user}_{now.strftime('%Y%m%d_%H%M%S')}.docx"
    out_path = os.path.join(get_handovers_dir(), filename)
    tpl.save(out_path)
    return out_path


def create_handover(data: dict, created_by: str = "") -> tuple[int, str]:
    """Logs the form to handover_records and renders its .docx, as one unit:
    the row is inserted first (the QR needs its id), the file rendered, then
    the row's docx_path filled in and only then committed - a rendering
    failure rolls the row back, so the history never lists a form with no
    file. Returns (record id, .docx path)."""
    conn = get_connection()
    try:
        now = dt.datetime.now().isoformat(timespec="seconds")
        ho_date = resolve_ho_date(data.get("ho_date")).isoformat()
        cur = conn.execute(
            """
            INSERT INTO handover_records (
                created_at, ho_date, user_no, user_name, branch_no, ho_type, reason, assets_json,
                ict_rep_name, ict_rep_id, receiving_name, receiving_title, receiving_dept,
                receiving_id, signature_receiving_name, signature_prepared_by, docx_path, created_by
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                now,
                ho_date,
                data.get("user_no", ""),
                data.get("receiving_name", ""),
                data.get("branch_no", ""),
                data.get("ho_type", ""),
                data.get("reason", ""),
                json.dumps(data.get("assets", [])),
                data.get("ict_rep_name", ""),
                data.get("ict_rep_id", ""),
                data.get("receiving_name", ""),
                data.get("receiving_title", ""),
                data.get("receiving_dept", ""),
                data.get("receiving_id", ""),
                data.get("signature_receiving_name", ""),
                data.get("signature_prepared_by", ""),
                "",
                created_by,
            ),
        )
        record_id = cur.lastrowid
        docx_path = render_handover_docx(data, record_id)
        conn.execute(
            "UPDATE handover_records SET docx_path = ? WHERE id = ?", (docx_path, record_id)
        )
        conn.commit()
        return record_id, docx_path
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def apply_handover_date(asset_ids: list[int], ho_date: str | None, performed_by: str = "") -> None:
    """Stamps asset_items.handover_date with the hand-over form's own date
    for every asset actually included on that form, once the form is
    generated - so Manage Assets reflects when the equipment was really
    handed over, rather than whatever date (often blank) an import last
    happened to carry for that row. Mirrors asset_edit.py's edit route:
    only rows whose handover_date is actually changing get updated/logged.
    Stored as dd/mm/yyyy, matching text_utils.normalize_handover_date's
    convention for this column (see importer.py/asset_edit.py) - not
    isoformat(), which is only used for handover_records.ho_date (a
    separate table, range-filtered by queries.py, so it must stay
    lexically sortable)."""
    new_value = resolve_ho_date(ho_date).strftime("%d/%m/%Y")
    conn = get_connection()
    try:
        for asset_id in asset_ids:
            row = conn.execute(
                "SELECT handover_date FROM asset_items WHERE id = ?", (asset_id,)
            ).fetchone()
            if row is None or (row["handover_date"] or "") == new_value:
                continue
            conn.execute(
                "UPDATE asset_items SET handover_date = ? WHERE id = ?", (new_value, asset_id)
            )
            log_activity(
                conn, "asset", "Hand-over form generated", performed_by=performed_by,
                target=f"Asset #{asset_id}", field="handover_date",
                old_value=row["handover_date"] or "", new_value=new_value,
            )
        conn.commit()
    finally:
        conn.close()
