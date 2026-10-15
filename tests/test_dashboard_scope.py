# -*- coding: utf-8 -*-
"""The Asset and CCTV Dashboards' All / Branches / Head Office switch
(?scope=) and Settings > Unit Groups - see queries.unit_group_sql. Same
isolated-DB pattern as tests/test_handover.py."""
import io
import os
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import create_app  # noqa: E402
from app.auth import create_account  # noqa: E402
from app.db import get_connection  # noqa: E402
from app.queries import get_current_asset_count, unit_group_sql  # noqa: E402

# branch_no -> (name, number of current assets)
UNITS = {
    "001": ("HANOI BRANCH", 3),
    "002": ("TAN BINH TRANSACTION OFFICE", 2),
    "003": ("ICT PLANNING DEPARTMENT", 4),
    "004": ("BUSINESS CENTER", 1),
    "C001": ("SCMC PICO", 5),
}


def _fresh_app():
    os.environ["LOCALAPPDATA"] = tempfile.mkdtemp(prefix="am_scopetest_")
    return create_app()


class DashboardScopeTests(unittest.TestCase):
    def setUp(self):
        self.app = _fresh_app()
        conn = get_connection()
        try:
            admin_role_id = conn.execute("SELECT id FROM roles WHERE name = 'Admin'").fetchone()["id"]
            conn.execute(
                "INSERT INTO import_batches (id, imported_at, kind, period) VALUES "
                "(1, datetime('now'), 'asset_report', '2026-01'), (2, datetime('now'), 'cctv_report', '2026-01')"
            )
            n = 0
            for branch_no, (name, count) in UNITS.items():
                conn.execute("INSERT INTO branches (branch_no, eng_name) VALUES (?, ?)", (branch_no, name))
                for _ in range(count):
                    n += 1
                    conn.execute(
                        "INSERT INTO asset_items (batch_id, asset_key, branch_no, branch_dept, device_name, "
                        "device_name_raw, serial_tag, status) VALUES (1, ?, ?, ?, 'PC', 'PC', ?, 'IN USE')",
                        (f"k{n}", branch_no, name, f"SN{n}"),
                    )
            # An unresolved row: only ever counted under "All".
            conn.execute(
                "INSERT INTO asset_items (batch_id, asset_key, branch_no, branch_dept, device_name, "
                "device_name_raw, serial_tag, status) VALUES (1, 'kx', '', 'SOMEWHERE', 'PC', 'PC', 'SNX', 'IN USE')"
            )
            conn.execute(
                "INSERT INTO cctv_items (batch_id, asset_key, branch_no, branch_dept, device_name, camera_count) "
                "VALUES (2, 'c1', '001', 'HANOI BRANCH', 'CCTV RECORDING', '8'), "
                "(2, 'c2', '003', 'ICT PLANNING DEPARTMENT', 'CCTV RECORDING', '4')"
            )
            conn.commit()
        finally:
            conn.close()
        create_account("scopetester", "scopetesterpass1", admin_role_id)
        self.client = self.app.test_client()
        self.client.post("/login", data={"username": "scopetester", "password": "scopetesterpass1"})

    def _groups(self):
        conn = get_connection()
        try:
            return {
                r["branch_no"]: r["g"]
                for r in conn.execute(f"SELECT branch_no, {unit_group_sql('b')} AS g FROM branches b")
            }
        finally:
            conn.close()

    def _counts(self):
        conn = get_connection()
        try:
            return {s: get_current_asset_count(conn, scope=s) for s in ("", "branch", "ho")}
        finally:
            conn.close()

    def test_auto_groups_by_name(self):
        self.assertEqual(
            self._groups(), {"001": "BRANCH", "002": "BRANCH", "003": "HO", "004": "HO", "C001": "HO"}
        )

    def test_counts_split_by_scope(self):
        # 3+2 branch network, 4+1+5 head office, +1 unresolved only under All.
        self.assertEqual(self._counts(), {"": 16, "branch": 5, "ho": 10})

    def test_unknown_scope_means_all(self):
        conn = get_connection()
        try:
            self.assertEqual(get_current_asset_count(conn, scope="bogus' OR 1=1 --"), 16)
        finally:
            conn.close()

    def test_dashboard_pages_render_per_scope(self):
        for url in ("/", "/?scope=branch", "/?scope=ho"):
            resp = self.client.get(url)
            self.assertEqual(resp.status_code, 200, url)
        body = self.client.get("/?scope=ho").get_data(as_text=True)
        self.assertIn("Asset Dashboard - Head Office", body)
        self.assertIn("ICT PLANNING DEPARTMENT", body)
        self.assertNotIn("HANOI BRANCH", body)

        body = self.client.get("/cctv/dashboard/?scope=branch").get_data(as_text=True)
        self.assertIn("HANOI BRANCH", body)
        self.assertNotIn("ICT PLANNING DEPARTMENT", body)

    def test_exports_follow_scope(self):
        resp = self.client.get("/export?scope=branch")
        self.assertIn("_branch.xlsx", resp.headers["Content-Disposition"])
        ws = openpyxl.load_workbook(io.BytesIO(resp.data)).worksheets[0]
        labels = {ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1)}
        self.assertEqual(labels, {"HANOI BRANCH", "TAN BINH TRANSACTION OFFICE"})

        resp = self.client.get("/cctv/dashboard/export-compare?scope=ho")
        ws = openpyxl.load_workbook(io.BytesIO(resp.data)).worksheets[0]
        labels = {ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1)}
        self.assertEqual(labels, {"ICT PLANNING DEPARTMENT"})

    def test_settings_override_moves_unit_and_logs(self):
        resp = self.client.post(
            "/settings/unit-groups",
            data={"group_004": "BRANCH", "group_001": "", "group_C001": "HO"},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(self._groups()["004"], "BRANCH")
        self.assertEqual(self._counts(), {"": 16, "branch": 6, "ho": 9})

        conn = get_connection()
        try:
            stored = {r["branch_no"]: r["unit_group"] for r in conn.execute("SELECT branch_no, unit_group FROM branches")}
            logs = conn.execute(
                "SELECT target, old_value, new_value FROM activity_log WHERE field = 'unit_group' ORDER BY id"
            ).fetchall()
        finally:
            conn.close()
        # C001 was already HO by name, but picking it explicitly pins it.
        self.assertEqual(stored["004"], "BRANCH")
        self.assertEqual(stored["C001"], "HO")
        self.assertIsNone(stored["001"])
        self.assertEqual([tuple(r) for r in logs], [("004", "AUTO", "BRANCH"), ("C001", "AUTO", "HO")])

        # Back to Auto.
        self.client.post("/settings/unit-groups", data={"group_004": ""})
        self.assertEqual(self._groups()["004"], "HO")


if __name__ == "__main__":
    unittest.main()
