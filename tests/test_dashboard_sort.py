# -*- coding: utf-8 -*-
"""Dashboard "Assets by Branch (by month)" Branch / Dept sort (?sort=branch /
-branch), on screen and in the Excel export. Same isolated-DB pattern as
tests/test_exports.py."""
import io
import os
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import create_app  # noqa: E402
from app.auth import create_account  # noqa: E402
from app.db import get_connection  # noqa: E402
from app.importer import import_asset_report  # noqa: E402
from app.routes.dashboard import _sort_month_table  # noqa: E402


def _row(branch_no, label, count):
    return {"branch_no": branch_no, "label": label, "cells": [], "current_count": count}


class SortFunctionTests(unittest.TestCase):
    ROWS = [_row("1", "ZETA BRANCH", 9), _row("", "Unresolved", 5), _row("2", "alpha branch", 3),
            _row("3", "MIDDLE BRANCH", 1)]

    def test_default_keeps_given_order(self):
        self.assertEqual([r["label"] for r in _sort_month_table(self.ROWS, "")],
                         ["ZETA BRANCH", "Unresolved", "alpha branch", "MIDDLE BRANCH"])

    def test_a_to_z_ignores_case_and_keeps_unresolved_last(self):
        self.assertEqual([r["label"] for r in _sort_month_table(self.ROWS, "branch")],
                         ["alpha branch", "MIDDLE BRANCH", "ZETA BRANCH", "Unresolved"])

    def test_z_to_a_keeps_unresolved_last(self):
        self.assertEqual([r["label"] for r in _sort_month_table(self.ROWS, "-branch")],
                         ["ZETA BRANCH", "MIDDLE BRANCH", "alpha branch", "Unresolved"])


class DashboardSortRouteTests(unittest.TestCase):
    def setUp(self):
        os.environ["LOCALAPPDATA"] = tempfile.mkdtemp(prefix="am_dashsort_")
        self.app = create_app()
        self.client = self.app.test_client()
        tmp = tempfile.mkdtemp(prefix="am_dashsort_files_")
        conn = get_connection()
        try:
            for no, name in (("001", "ZETA BRANCH"), ("002", "ALPHA BRANCH")):
                conn.execute(
                    "INSERT INTO branches (branch_no, local_name, eng_name, updated_at) VALUES (?, ?, ?, datetime('now'))",
                    (no, name, name),
                )
            admin = conn.execute("SELECT id FROM roles WHERE name = 'Admin'").fetchone()["id"]
            conn.commit()
        finally:
            conn.close()
        # ZETA has more assets, so it's first in the default (count) order.
        for label, n in (("Zeta Branch", 3), ("Alpha Branch", 1)):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append([None] * 7 + ["Branch/TO/Center Name:", label])
            ws.append(["NO", "BRANCH / DEPT", "DEVICE NAME", "USER ID", "FULL NAME", "IP", "MODEL DEVICE",
                       "SERIAL/ SERVICE TAG", "STATUS", "REMARK"])
            for i in range(n):
                ws.append([i + 1, label.upper(), "PC", f"{i}", "A", "", "DELL", f"SN-{label[:1]}-{i}", "USING LOCAL", ""])
            path = os.path.join(tmp, f"{label}.xlsx")
            wb.save(path)
            import_asset_report(path, source_label=os.path.basename(path), period="2026-03")
        create_account("sorter", "sortpass1", admin)
        self.client.post("/login", data={"username": "sorter", "password": "sortpass1"})

    def _order(self, html: bytes):
        """The two branch names in the order they first appear on the page."""
        return sorted(("ZETA BRANCH", "ALPHA BRANCH"), key=lambda n: html.index(n.encode()))

    def test_default_then_a_to_z_then_z_to_a(self):
        page = self.client.get("/?year=2026").data
        self.assertEqual(self._order(page), ["ZETA BRANCH", "ALPHA BRANCH"])
        self.assertIn(b"sort=branch", page)  # header links to A-Z next

        page = self.client.get("/?year=2026&sort=branch").data
        self.assertEqual(self._order(page), ["ALPHA BRANCH", "ZETA BRANCH"])
        self.assertIn(b'name="sort" value="branch"', page)  # year switch keeps the sort
        self.assertIn(b"sort=-branch", page)

        page = self.client.get("/?year=2026&sort=-branch").data
        self.assertEqual(self._order(page), ["ZETA BRANCH", "ALPHA BRANCH"])

    def test_unknown_sort_is_default(self):
        page = self.client.get("/?year=2026&sort=bogus").data
        self.assertEqual(self._order(page), ["ZETA BRANCH", "ALPHA BRANCH"])

    def test_export_follows_the_sort(self):
        resp = self.client.get("/export?year=2026&sort=branch")
        ws = openpyxl.load_workbook(io.BytesIO(resp.data)).worksheets[0]
        labels = [row[0] for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]
        self.assertEqual(labels[:2], ["ALPHA BRANCH", "ZETA BRANCH"])


if __name__ == "__main__":
    unittest.main()
