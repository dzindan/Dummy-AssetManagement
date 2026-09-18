# -*- coding: utf-8 -*-
"""CCTV Dashboard rendering, including the per-branch compare/expand tree.

Follows the project's isolated-DB pattern (see tests/test_auth.py) and
actually renders the page through app.test_client() rather than just
calling the route function directly - a Jinja error (e.g. a dict key that
collides with a builtin method name, like `items`) only ever surfaces at
render time, not from the view function's own Python.
"""
import io
import os
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import create_app  # noqa: E402
from app.importer import import_asset_report  # noqa: E402


def _fresh_app():
    os.environ["LOCALAPPDATA"] = tempfile.mkdtemp(prefix="am_cctvdashtest_")
    return create_app()


def _build_cctv_workbook(path: str, branch_name: str) -> None:
    wb = openpyxl.Workbook()
    cctv = wb.active
    cctv.title = "CCTV REPORT"
    cctv.append([None] * 10 + ["Branch/TO/Center Name:", branch_name])
    cctv.append(
        ["NO", "BRANCH / DEPT", "DEVICE NAME", "IP ADDRESS", "PRODUCTION", "MODEL DEVICE", "SERIAL NO",
         "STATUS", "NUMBER  OF CAMERA CONNECTED", "NUMBER OF HARD DISK", "CAPACITY OF ALL HARD DISK",
         "LOCATION", "REMARK"]
    )
    cctv.append(
        [1, branch_name, "CCTV RECORDING 1", "10.0.1.1", "HIK VISION", "DS-7316", "SN-CCTV-1",
         "USING LOCAL", 16, "3", "24TB", "IT ROOM", ""]
    )
    cctv.append(
        [2, branch_name, "CCTV RECORDING 2", "10.0.1.2", "HIK VISION", "DS-7324", "SN-CCTV-2",
         "USING LOCAL", "8 (Live 7)", "2 HDD", "16TB (2X8TB) Total", "IT ROOM", ""]
    )
    wb.save(path)


class CctvDashboardRenderTests(unittest.TestCase):
    def setUp(self):
        self.app = _fresh_app()
        self.client = self.app.test_client()
        resp = self.client.post(
            "/setup",
            data={"username": "admin", "password": "adminpass123", "confirm": "adminpass123"},
            follow_redirects=False,
        )
        self.assertEqual(resp.status_code, 302)
        self.tmpdir = tempfile.mkdtemp(prefix="am_cctvdashtest_import_")

    def test_dashboard_renders_with_branch_compare_tree(self):
        path = os.path.join(self.tmpdir, "report.xlsx")
        _build_cctv_workbook(path, "Test Branch")
        reports = import_asset_report(path, source_label="report.xlsx", period="2026-03")
        self.assertEqual(reports[0].error, "")

        resp = self.client.get("/cctv/dashboard/")
        self.assertEqual(resp.status_code, 200)
        body = resp.data.decode("utf-8")
        self.assertIn("Compare by Branch", body)
        self.assertIn("2</strong> DVR/Recorder", body)
        # 16 + 8 (the "8 (Live 7)" row's leading int) = 24 cameras.
        self.assertIn("24</strong> Cameras", body)
        # 3 + 2 ("2 HDD" row's leading int) = 5.
        self.assertIn("5</strong> HDD Count", body)
        # 24TB + 16TB (from "16TB (2X8TB) Total") = 40.00 TB.
        self.assertIn("40.00 TB", body)

        # Month-by-month breakdown for the report's own period (2026-03) -
        # same underlying numbers as the branch-level totals above, but
        # this is the actual thing that was missing before: per-branch
        # comparison "of each month", not just the current snapshot.
        self.assertIn("By Month (2026)", body)
        self.assertIn("2026-03", body)
        self.assertIn("Current Items", body)

    def test_by_month_table_shows_delta_vs_previous_period(self):
        """The main Asset Dashboard's by-month table shows +added/-removed
        against the previous period on record for every column - the
        per-branch tree's own By Month table needs the same, not just raw
        counts with nothing to compare them to."""
        march_path = os.path.join(self.tmpdir, "march.xlsx")
        _build_cctv_workbook(march_path, "Test Branch")
        reports = import_asset_report(march_path, source_label="march.xlsx", period="2026-03")
        self.assertEqual(reports[0].error, "")

        april = os.path.join(self.tmpdir, "april.xlsx")
        wb = openpyxl.Workbook()
        cctv = wb.active
        cctv.title = "CCTV REPORT"
        cctv.append([None] * 10 + ["Branch/TO/Center Name:", "Test Branch"])
        cctv.append(
            ["NO", "BRANCH / DEPT", "DEVICE NAME", "IP ADDRESS", "PRODUCTION", "MODEL DEVICE", "SERIAL NO",
             "STATUS", "NUMBER  OF CAMERA CONNECTED", "NUMBER OF HARD DISK", "CAPACITY OF ALL HARD DISK",
             "LOCATION", "REMARK"]
        )
        cctv.append(
            [1, "Test Branch", "CCTV RECORDING 1", "10.0.1.1", "HIK VISION", "DS-7316", "SN-CCTV-1",
             "USING LOCAL", 16, "3", "24TB", "IT ROOM", ""]
        )
        wb.save(april)
        reports = import_asset_report(april, source_label="april.xlsx", period="2026-04")
        self.assertEqual(reports[0].error, "")

        resp = self.client.get("/cctv/dashboard/")
        self.assertEqual(resp.status_code, 200)
        body = resp.data.decode("utf-8")

        # April (1 DVR/16 cameras/3 HDD/24TB) vs March (2/24/5/40.00) -
        # every metric should show its own negative delta.
        self.assertIn('<span style="color:var(--danger);">-1</span>', body)
        self.assertIn('<span style="color:var(--danger);">-8</span>', body)
        self.assertIn('<span style="color:var(--danger);">-2</span>', body)
        self.assertIn('<span style="color:var(--danger);">-16.00</span>', body)

    def test_export_compare_has_current_and_by_month_sheets(self):
        path = os.path.join(self.tmpdir, "report.xlsx")
        _build_cctv_workbook(path, "Test Branch")
        reports = import_asset_report(path, source_label="report.xlsx", period="2026-03")
        self.assertEqual(reports[0].error, "")

        resp = self.client.get("/cctv/dashboard/export-compare?year=2026")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp.content_type,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        wb = openpyxl.load_workbook(io.BytesIO(resp.data))
        self.assertIn("Current by Branch", wb.sheetnames)
        self.assertIn("By Month 2026", wb.sheetnames)

        current = wb["Current by Branch"]
        self.assertEqual(
            [c.value for c in current[1]],
            ["Branch", "DVR/Recorder", "Cameras", "HDD Count", "HDD Capacity (TB)"],
        )
        data_row = [c.value for c in current[2]]
        self.assertEqual(data_row, ["Test Branch", 2, 24, 5, 40.0])

        month_ws = wb["By Month 2026"]
        month_row = [c.value for c in month_ws[2]]
        self.assertEqual(month_row[:4], ["Test Branch", "2026-03", 2, None])  # no prior period -> delta None

    def test_dashboard_renders_with_no_cctv_data(self):
        resp = self.client.get("/cctv/dashboard/")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("No CCTV items currently on file.", resp.data.decode("utf-8"))


if __name__ == "__main__":
    unittest.main()
