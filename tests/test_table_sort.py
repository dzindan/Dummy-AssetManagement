"""Server-side column sort for Manage Assets / Manage CCTV (app/sorting.py,
queries.ASSET_SORTS / CCTV_SORTS): whitelisted keys only, blanks last in
both directions, dd/mm/yyyy dates in date order, free-text counts in number
order, and the header links / filter form / export keep the sort.
Fixture pattern from tests/test_hand_fix.py."""
import io
import os
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import create_app  # noqa: E402
from app.auth import create_account  # noqa: E402
from app.db import get_connection  # noqa: E402
from app.importer import import_asset_report  # noqa: E402
from app.queries import ASSET_SORTS, order_by_sql, search_assets, search_cctv  # noqa: E402
from app.sorting import next_sort  # noqa: E402


def _workbook(path: str) -> None:
    wb = openpyxl.Workbook()
    oa = wb.active
    oa.title = "OA EQUIPMENT"
    oa.append([None] * 7 + ["Branch/TO/Center Name:", "Test Branch"])
    oa.append(["NO", "BRANCH / DEPT", "DEVICE NAME", "USER ID", "FULL NAME", "IP", "MODEL DEVICE",
               "SERIAL/ SERVICE TAG", "STATUS", "REMARK", "HANDOVER DATE"])
    oa.append([1, "TEST BRANCH", "PC", "1001", "BINH", "", "DELL", "SN-1", "USING LOCAL", "", "13/01/2020"])
    oa.append([2, "TEST BRANCH", "PC", "1002", "AN", "", "DELL", "SN-2", "USING LOCAL", "", "01/08/2016"])
    oa.append([3, "TEST BRANCH", "PC", "1003", "", "", "DELL", "SN-3", "USING LOCAL", "", "NA"])
    oa.append([4, "TEST BRANCH", "PC", "1004", "CUONG", "", "DELL", "SN-4", "USING LOCAL", "", "20/07/2019"])
    cctv = wb.create_sheet("CCTV REPORT")
    cctv.append([None] * 10 + ["Branch/TO/Center Name:", "Test Branch"])
    cctv.append(["NO", "BRANCH / DEPT", "DEVICE NAME", "IP ADDRESS", "PRODUCTION", "MODEL DEVICE", "SERIAL NO",
                 "STATUS", "NUMBER  OF CAMERA CONNECTED", "NUMBER OF HARD DISK", "CAPACITY OF ALL HARD DISK",
                 "LOCATION", "REMARK"])
    cctv.append([1, "TEST BRANCH", "DVR", "", "HIK", "A", "DVR-1", "USING LOCAL", 9, "1", "24TB", "", ""])
    cctv.append([2, "TEST BRANCH", "DVR", "", "HIK", "B", "DVR-2", "USING LOCAL", 14, "2", "7452.04 GB", "", ""])
    cctv.append([3, "TEST BRANCH", "DVR", "", "HIK", "C", "DVR-3", "USING LOCAL", "", "3", "", "", ""])
    wb.save(path)


class SortHelperTests(unittest.TestCase):
    def test_cycle(self):
        self.assertEqual(next_sort("", "status"), "status")
        self.assertEqual(next_sort("status", "status"), "-status")
        self.assertEqual(next_sort("-status", "status"), "")
        self.assertEqual(next_sort("-device", "status"), "status")

    def test_unknown_key_falls_back_to_default(self):
        self.assertEqual(order_by_sql("bk.id; DROP TABLE x", ASSET_SORTS, "DEFAULT"), "DEFAULT")
        self.assertEqual(order_by_sql("", ASSET_SORTS, "DEFAULT"), "DEFAULT")


class ServerSortTests(unittest.TestCase):
    def setUp(self):
        os.environ["LOCALAPPDATA"] = tempfile.mkdtemp(prefix="am_sort_")
        self.app = create_app()
        self.client = self.app.test_client()
        tmp = tempfile.mkdtemp(prefix="am_sort_files_")
        conn = get_connection()
        try:
            conn.execute("INSERT INTO branches (branch_no, local_name, eng_name, updated_at) "
                         "VALUES ('001', 'Test Branch', 'TEST BRANCH', datetime('now'))")
            admin = conn.execute("SELECT id FROM roles WHERE name = 'Admin'").fetchone()["id"]
            conn.commit()
        finally:
            conn.close()
        create_account("sorter", "sortpass123", admin)
        self.client.post("/login", data={"username": "sorter", "password": "sortpass123"})
        path = os.path.join(tmp, "2026-07.xlsx")
        _workbook(path)
        import_asset_report(path, source_label="2026-07.xlsx", period="2026-07")
        self.conn = get_connection()

    def tearDown(self):
        self.conn.close()

    def _assets(self, sort):
        rows, _ = search_assets(self.conn, {"device_name": ["PC"]}, sort=sort)
        return [r["serial_tag"] for r in rows]

    def test_text_sort_blanks_last_both_ways(self):
        self.assertEqual(self._assets("full_name"), ["SN-2", "SN-1", "SN-4", "SN-3"])
        self.assertEqual(self._assets("-full_name"), ["SN-4", "SN-1", "SN-2", "SN-3"])

    def test_handover_date_order(self):
        self.assertEqual(self._assets("handover"), ["SN-2", "SN-4", "SN-1", "SN-3"])  # 2016, 2019, 2020, NA
        self.assertEqual(self._assets("usage"), ["SN-1", "SN-4", "SN-2", "SN-3"])  # shortest use first

    def test_ip_column_sorts_numerically_and_exports(self):
        conn = self.conn
        for serial, ip in (("SN-1", "10.0.0.10"), ("SN-2", "10.0.0.9"), ("SN-3", "DHCP"), ("SN-4", "9.1.1.1")):
            conn.execute("UPDATE asset_items SET ip = ? WHERE serial_tag = ?", (ip, serial))
        conn.commit()
        self.assertEqual(self._assets("ip"), ["SN-4", "SN-2", "SN-1", "SN-3"])  # 9.x < 10.0.0.9 < 10.0.0.10, text last
        html = self.client.get("/assets/?device_name=PC").get_data(as_text=True)
        self.assertIn('data-field="ip" data-value="10.0.0.9"', html)
        resp = self.client.get("/assets/export?device_name=PC&sort=ip")
        ws = openpyxl.load_workbook(io.BytesIO(resp.data)).active
        header = [c.value for c in ws[1]]
        ips = [row[header.index("IP")] for row in ws.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(ips, ["9.1.1.1", "10.0.0.9", "10.0.0.10", "DHCP"])

    def test_cctv_number_sorts(self):
        def serials(sort):
            rows, _ = search_cctv(self.conn, {}, sort=sort)
            return [r["serial_tag"] for r in rows]
        self.assertEqual(serials("-cameras"), ["DVR-2", "DVR-1", "DVR-3"])  # 14, 9, blank
        self.assertEqual(serials("hdd_capacity"), ["DVR-2", "DVR-1", "DVR-3"])  # ~7.3 TB, 24 TB, blank

    def test_page_links_keep_filters_and_form_keeps_sort(self):
        html = self.client.get("/assets/?device_name=PC&sort=-full_name&page=2").get_data(as_text=True)
        self.assertIn('href="/assets/?device_name=PC&amp;sort=status"', html)  # another column: ascending
        self.assertIn('href="/assets/?device_name=PC"', html)  # this column: back to default
        self.assertIn('<input type="hidden" name="sort" value="-full_name">', html)
        self.assertLess(html.index("CUONG"), html.index("BINH"))

    def test_export_follows_sort(self):
        resp = self.client.get("/assets/export?device_name=PC&sort=-full_name")
        ws = openpyxl.load_workbook(io.BytesIO(resp.data)).active
        header = [c.value for c in ws[1]]
        names = [row[header.index("Full Name")] for row in ws.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(names[:3], ["CUONG", "BINH", "AN"])

    def test_cctv_page_renders_sorted(self):
        resp = self.client.get("/cctv/?sort=-cameras")
        self.assertEqual(resp.status_code, 200)
        html = resp.get_data(as_text=True)
        self.assertLess(html.index("DVR-2"), html.index("DVR-1"))


if __name__ == "__main__":
    unittest.main()
