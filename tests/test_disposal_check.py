# -*- coding: utf-8 -*-
"""Disposal Check: an Aither disposal / inventory list (two header lines
under merged cells, a second sheet whose items are also copied to the
first) checked item by item against current assets by serial. Same
isolated-DB pattern as tests/test_custom_branch.py."""
import io
import os
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import create_app  # noqa: E402
from app.auth import create_account  # noqa: E402
from app.db import get_connection  # noqa: E402

HEAD1 = ["NO", "c", "Branch ID", "Manage Branch ID", "FA/ WT Name", "Ref. Number", "Register Date", "FA/ WT Code",
         "Classification", "Serial number/ Specification", "Aither system", None, None, "Inventoried", None, None,
         "Difference", None, None, "Original Location", "Current Location", "Aither Status", "Inventory Status",
         "Remark/Suggestion", "Checker", "Approver", "Report status", "GAD opinion"]
HEAD2 = [None] * 10 + ["Quantity", "Historical Cost Price", "Book value"] * 3 + [None] * 9


def item(no, name, ref, cls, serial, status="20-Broken", remark="Disposal", gad=""):
    row = [no, "1", 8064, 8064, name, ref, "01/01/2020", "2110101-OTHERS", cls, serial] + [1, 0, 0] * 3
    return row + ["TDH", "TDH", "10-Normal", status, remark, "", "", "30-BU-Approved", gad]


class DisposalCheckTests(unittest.TestCase):
    def setUp(self):
        os.environ["LOCALAPPDATA"] = tempfile.mkdtemp(prefix="am_disposal_")
        self.app = create_app()
        self.client = self.app.test_client()
        conn = get_connection()
        try:
            conn.execute("INSERT INTO branches (branch_no, local_name, eng_name, updated_at) "
                         "VALUES ('8064', 'TRAN DUY HUNG', 'TRAN DUY HUNG BRANCH', datetime('now')), "
                         "('8017', 'HA NOI', 'HA NOI BRANCH', datetime('now'))")
            old = conn.execute("INSERT INTO import_batches (imported_at, kind, label, period) "
                               "VALUES (datetime('now'), 'asset_report', 'TDH', '2026-06')").lastrowid
            new = conn.execute("INSERT INTO import_batches (imported_at, kind, label, period) "
                               "VALUES (datetime('now'), 'asset_report', 'TDH', '2026-08')").lastrowid
            hn = conn.execute("INSERT INTO import_batches (imported_at, kind, label, period) "
                              "VALUES (datetime('now'), 'asset_report', 'HN', '2026-08')").lastrowid
            rows = [
                (new, "8064", "PC", "SN-OK", "WAREHOUSE", ""),
                (new, "8064", "PC", "SN-USED", "USING LOCAL", "19506211"),
                (hn, "8017", "LCD", "SN-HN", "WAREHOUSE", ""),
                (old, "8064", "PC", "SN-OLD", "BROKEN", ""),
            ]
            for i, (batch, branch, dev, serial, status, uid) in enumerate(rows):
                conn.execute("INSERT INTO asset_items (batch_id, asset_key, branch_dept, branch_no, device_name, "
                             "serial_tag, status, user_id_raw) VALUES (?, ?, 'x', ?, ?, ?, ?, ?)",
                             (batch, f"k{i}", branch, dev, serial, status, uid))
            admin = conn.execute("SELECT id FROM roles WHERE name = 'Admin'").fetchone()["id"]
            conn.commit()
        finally:
            conn.close()
        create_account("checker", "checkpass123", admin)
        self.client.post("/login", data={"username": "checker", "password": "checkpass123"})

    def _workbook(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "sheet1"
        ws.append(HEAD1)
        ws.append(HEAD2)
        # Aither's own export repeats the merged header text on line 2.
        ws.append(["", "", "", "", "FA/ WT Name", "Ref. Number", "", "", "", "Serial number/ Specification"])
        ws.append(item(1, "PC ok", "R1", "PC", "SN-OK"))
        ws.append(item(2, "PC still used", "R2", "PC", "sn ok-used".replace("ok-", "-").upper()))  # "SN -USED"
        ws.append(item(3, "Monitor at HN", "R3", "MONITOR", "SN-HN"))
        ws.append(item(4, "Old PC", "R4", "PC", "SN-OLD", gad="Bus make Memo to dispose in Aither"))
        ws.append(item(5, "Unknown", "R5", "PC", "SN-NEW"))
        ws.append(item(6, "Sofa", "R6", "Sofa", "N/A"))
        ws.append(item(7, "Printer kept", "R7", "Printer", "SN-NEW", status="10-Normal", remark=""))
        ws.append(item(8, "Desk kept", "R8", "Desk", "N/A", status="10-Normal", remark=""))
        ws2 = wb.create_sheet("sheet2")
        ws2.append(["NO", "c", "Branch ID", "Manage Branch ID", "Transfer Branch", "FA/ WT Name", "Ref. Number",
                    "Register Date", "FA/ WT Code", "Classification", "Serial number/ Specification"])
        ws2.append([1, "1", 8064, 8064, 8064, "PC ok", "R1", "07/10/2026", "x", "PC", "SN-OK"])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def _upload(self, only=True):
        data = {"file": (self._workbook(), "8064_list.xlsx")}
        if only:
            data["only_marked"] = "1"
        resp = self.client.post("/disposal-check/", data=data, content_type="multipart/form-data")
        self.assertEqual(resp.status_code, 302)
        return resp.headers["Location"]

    def test_results(self):
        from app.disposal import check_items, read_disposal_list
        from app.routes.disposal_check import _token_path
        url = self._upload()
        with self.app.test_request_context():
            token = url.split("/disposal-check/")[1].split("?")[0]
            items, notes = read_disposal_list(_token_path(token))
        self.assertEqual(len(items), 8)                     # sheet2's R1 merged into sheet1's R1
        self.assertEqual(items[0].also_in, ["sheet2 row 2"])
        conn = get_connection()
        try:
            check_items(conn, items)
        finally:
            conn.close()
        got = {i.ref: i.result for i in items}
        self.assertEqual(got, {"R1": "ok", "R2": "in_use", "R3": "other_branch", "R4": "not_current",
                               "R5": "duplicate", "R6": "no_serial", "R7": "duplicate", "R8": "no_serial"})
        self.assertIn("19506211", items[1].note)
        self.assertIn("Last seen 2026-06 at 8064", items[3].note)
        self.assertIn("R7 Printer kept", items[4].note)
        self.assertEqual([i.ref for i in items if i.marked_for_disposal], ["R1", "R2", "R3", "R4", "R5", "R6"])

    def test_page_filter_and_export(self):
        url = self._upload()
        page = self.client.get(url).data.decode()
        self.assertIn("6 item(s) marked for disposal (of 8 on the list)", page)
        self.assertIn("Still in use: 1", page)
        self.assertIn("SN-USED", page.replace(" ", ""))
        self.assertNotIn("<td>Printer kept</td>", page)
        only_used = self.client.get(url + "&result=in_use").data.decode()
        self.assertIn("PC still used", only_used)
        self.assertNotIn("Monitor at HN", only_used)
        token = url.split("/disposal-check/")[1].split("?")[0]
        resp = self.client.get(f"/disposal-check/{token}/export.xlsx?only=0")
        self.assertEqual(resp.status_code, 200)
        ws = openpyxl.load_workbook(io.BytesIO(resp.data)).active
        self.assertEqual(ws.cell(1, 1).value, "Result")
        self.assertEqual(ws.max_row, 9)                     # header + all 8 items

    def test_bad_token_and_file_type(self):
        self.assertEqual(self.client.get("/disposal-check/../../x.xlsx").status_code, 404)
        self.assertEqual(self.client.get("/disposal-check/" + "a" * 32 + ".xlsx").status_code, 404)
        resp = self.client.post("/disposal-check/", data={"file": (io.BytesIO(b"x"), "list.csv")},
                                content_type="multipart/form-data", follow_redirects=True)
        self.assertIn(b"Choose the disposal list", resp.data)


if __name__ == "__main__":
    unittest.main()
