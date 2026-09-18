# -*- coding: utf-8 -*-
"""Branch Detail page rendering, including its CCTV section.

Follows the project's isolated-DB pattern (see tests/test_auth.py) and
actually renders the page through app.test_client() - a Jinja error only
ever surfaces at render time (see tests/test_cctv_dashboard.py's own
docstring on the exact bug that caught).
"""
import io
import os
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import create_app  # noqa: E402
from app.db import get_connection  # noqa: E402
from app.importer import import_asset_report  # noqa: E402


def _fresh_app():
    os.environ["LOCALAPPDATA"] = tempfile.mkdtemp(prefix="am_branchdetailtest_")
    return create_app()


def _build_two_sheet_workbook(path: str, branch_name: str) -> None:
    import openpyxl as _ox

    wb = _ox.Workbook()
    oa = wb.active
    oa.title = "OA EQUIPMENT"
    oa.append([None] * 7 + ["Branch/TO/Center Name:", branch_name])
    oa.append(
        ["NO", "BRANCH / DEPT", "DEVICE NAME", "USER ID", "FULL NAME", "IP", "MODEL DEVICE",
         "SERIAL/ SERVICE TAG", "STATUS", "REMARK"]
    )
    oa.append([1, branch_name, "PC", "1001", "NGUYEN VAN A", "10.0.0.1", "DELL 3060", "SN-PC-1", "USING LOCAL", ""])

    cctv = wb.create_sheet("CCTV REPORT")
    cctv.append([None] * 10 + ["Branch/TO/Center Name:", branch_name])
    cctv.append(
        ["NO", "BRANCH / DEPT", "DEVICE NAME", "IP ADDRESS", "PRODUCTION", "MODEL DEVICE", "SERIAL NO",
         "STATUS", "NUMBER  OF CAMERA CONNECTED", "NUMBER OF HARD DISK", "CAPACITY OF ALL HARD DISK",
         "LOCATION", "REMARK"]
    )
    cctv.append(
        [1, branch_name, "CCTV RECORDING 1", "10.0.1.1", "HIK VISION", "DS-7316", "SN-CCTV-1",
         "USING LOCAL", 16, "3", "24TB", "IT ROOM", ""]
    )
    wb.save(path)


class BranchDetailCctvSectionTests(unittest.TestCase):
    def setUp(self):
        self.app = _fresh_app()
        self.client = self.app.test_client()
        resp = self.client.post(
            "/setup",
            data={"username": "admin", "password": "adminpass123", "confirm": "adminpass123"},
            follow_redirects=False,
        )
        self.assertEqual(resp.status_code, 302)

        conn = get_connection()
        try:
            conn.execute(
                "INSERT INTO branches (branch_no, local_name, eng_name, updated_at) "
                "VALUES ('001', 'Test Branch', 'TEST BRANCH', datetime('now'))"
            )
            conn.commit()
        finally:
            conn.close()

        self.tmpdir = tempfile.mkdtemp(prefix="am_branchdetailtest_import_")
        path = os.path.join(self.tmpdir, "report.xlsx")
        _build_two_sheet_workbook(path, "Test Branch")
        reports = import_asset_report(path, source_label="report.xlsx", period="2026-03")
        for r in reports:
            self.assertEqual(r.error, "")

    def test_detail_page_renders_cctv_section(self):
        resp = self.client.get("/branch/001")
        self.assertEqual(resp.status_code, 200)
        body = resp.data.decode("utf-8")

        self.assertIn("CCTV Item Count Trend", body)
        self.assertIn("Current CCTV Items", body)
        self.assertIn("CCTV RECORDING 1", body)
        # 1 recorder, 16 cameras, 3 HDDs, 24.00 TB - the stat tiles.
        self.assertIn("1</div>\n    <div class=\"label\">DVR/Recorder", body)

    def test_export_cctv_route(self):
        resp = self.client.get("/branch/001/export-cctv")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp.content_type,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        wb = openpyxl.load_workbook(io.BytesIO(resp.data))
        self.assertIn("Current CCTV", wb.sheetnames)
        ws = wb["Current CCTV"]
        row = [c.value for c in ws[2]]
        self.assertEqual(row[0], "CCTV RECORDING 1")
        self.assertEqual(row[4], "SN-CCTV-1")

    def test_detail_page_renders_with_no_cctv_data(self):
        conn = get_connection()
        try:
            conn.execute(
                "INSERT INTO branches (branch_no, local_name, eng_name, updated_at) "
                "VALUES ('002', 'Empty Branch', 'EMPTY BRANCH', datetime('now'))"
            )
            conn.commit()
        finally:
            conn.close()
        resp = self.client.get("/branch/002")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("No CCTV items on record for this branch.", resp.data.decode("utf-8"))


if __name__ == "__main__":
    unittest.main()
