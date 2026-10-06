"""CCTV and DVR devices in asset_items are left out of the Asset Dashboard /
Branch Detail counts and charts and shown on their own, while listings still
return every row. Anything else - even a CCTV sheet's monitor (an LCD) -
stays an asset. Fixture pattern from tests/test_hand_fix.py.
"""
import io
import os
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import analytics, create_app  # noqa: E402
from app.auth import create_account  # noqa: E402
from app.db import get_connection  # noqa: E402
from app.importer import import_asset_report  # noqa: E402
from app.queries import get_current_asset_count, get_current_assets, is_cctv_device  # noqa: E402


def _workbook(path: str) -> None:
    """OA sheet: a PC, an office LCD and a CCTV device. CCTV sheet: a DVR
    and its monitor (an LCD) - 3 assets (PC + 2 LCDs), 2 CCTV/DVR."""
    wb = openpyxl.Workbook()
    oa = wb.active
    oa.title = "OA EQUIPMENT"
    oa.append([None] * 7 + ["Branch/TO/Center Name:", "Test Branch"])
    oa.append(["NO", "BRANCH / DEPT", "DEVICE NAME", "USER ID", "FULL NAME", "IP", "MODEL DEVICE",
               "SERIAL/ SERVICE TAG", "STATUS", "REMARK"])
    oa.append([1, "TEST BRANCH", "PC", "1001", "NGUYEN VAN A", "", "DELL 3060", "SN-PC", "USING LOCAL", ""])
    oa.append([2, "TEST BRANCH", "LCD", "1001", "NGUYEN VAN A", "", "DELL P2219", "SN-LCD", "USING LOCAL", ""])
    oa.append([3, "TEST BRANCH", "CCTV", "", "", "", "HIKVISION", "SN-CAM", "USING LOCAL", ""])
    cctv = wb.create_sheet("CCTV REPORT")
    cctv.append([None] * 10 + ["Branch/TO/Center Name:", "Test Branch"])
    cctv.append(["NO", "BRANCH / DEPT", "DEVICE NAME", "IP ADDRESS", "PRODUCTION", "MODEL DEVICE", "SERIAL NO",
                 "STATUS", "NUMBER  OF CAMERA CONNECTED", "NUMBER OF HARD DISK", "CAPACITY OF ALL HARD DISK",
                 "LOCATION", "REMARK"])
    cctv.append([1, "TEST BRANCH", "DVR/CCTV RECORDER", "", "HIKVISION", "DS-7616", "SN-DVR", "USING LOCAL",
                 8, "2", "8TB", "SERVER ROOM", ""])
    cctv.append([2, "TEST BRANCH", "LCD", "", "SAMSUNG", "S22", "SN-MON", "USING LOCAL",
                 "", "", "", "SERVER ROOM", ""])
    wb.save(path)


class CctvGearSplitTests(unittest.TestCase):
    def setUp(self):
        os.environ["LOCALAPPDATA"] = tempfile.mkdtemp(prefix="am_cctvsplit_")
        self.app = create_app()
        self.client = self.app.test_client()
        tmp = tempfile.mkdtemp(prefix="am_cctvsplit_files_")
        conn = get_connection()
        try:
            conn.execute("INSERT INTO branches (branch_no, local_name, eng_name, updated_at) "
                         "VALUES ('001', 'Test Branch', 'TEST BRANCH', datetime('now'))")
            admin = conn.execute("SELECT id FROM roles WHERE name = 'Admin'").fetchone()["id"]
            conn.commit()
        finally:
            conn.close()
        create_account("viewer", "viewpass123", admin)
        self.client.post("/login", data={"username": "viewer", "password": "viewpass123"})
        for period in ("2026-06", "2026-07"):
            path = os.path.join(tmp, f"{period}.xlsx")
            _workbook(path)
            import_asset_report(path, source_label=f"{period}.xlsx", period=period)
        self.conn = get_connection()

    def tearDown(self):
        self.conn.close()

    def test_fixture_has_all_five_rows_listed(self):
        rows = get_current_assets(self.conn, branch_no="001")
        self.assertEqual(len(rows), 5)
        self.assertEqual(len([r for r in rows if not is_cctv_device(r["device_name"])]), 3)

    def test_dashboard_counts(self):
        self.assertEqual(get_current_asset_count(self.conn), 3)
        self.assertEqual(get_current_asset_count(self.conn, cctv_gear=True), 2)

    def test_trends_leave_cctv_gear_out(self):
        _p, items, matrix = analytics.get_all_branches_item_trend(self.conn, max_series=None)
        self.assertEqual(sorted(items), ["LCD", "PC"])
        self.assertEqual(matrix["LCD"], {"2026-06": 2, "2026-07": 2})  # the CCTV monitor still counts
        _p, items, _m = analytics.get_branch_item_trend(self.conn, "001", max_series=None)
        self.assertEqual(sorted(items), ["LCD", "PC"])

    def test_cctv_tables_unaffected(self):
        _p, items, _m = analytics.get_all_branches_item_trend(self.conn, table="cctv_items")
        self.assertEqual(sorted(items), ["DVR/CCTV RECORDER", "LCD"])

    def test_month_year_tables(self):
        _p, _t, totals, _a, _r = analytics.get_branch_month_change_table(self.conn, "2026")
        self.assertEqual(totals[5:7], [3, 3])
        _p, _t, totals, _a, _r = analytics.get_branch_device_year_table(self.conn, "001", "2026", ["LCD", "PC"])
        self.assertEqual(totals[5:7], [3, 3])
        self.assertEqual(analytics.get_year_comparison_table(self.conn), [{"year": "2026", "count": 3, "change": None}])

    def test_device_type_table(self):
        _p, table, totals, _a, _r = analytics.get_device_month_change_table(self.conn, "2026")
        self.assertEqual({r["item"]: r["cells"][6]["count"] for r in table}, {"LCD": 2, "PC": 1})
        # Same totals as the by-branch table, month by month.
        _p, _t, branch_totals, _a, _r = analytics.get_branch_month_change_table(self.conn, "2026")
        self.assertEqual(totals, branch_totals)
        # July vs. June: same assets, so no movement.
        self.assertEqual((table[0]["cells"][6]["added"], table[0]["cells"][6]["removed"]), (0, 0))

    def test_device_type_table_on_page_and_export(self):
        html = self.client.get("/").get_data(as_text=True)
        self.assertIn("Assets by Device Type (by month)", html)
        resp = self.client.get("/export?year=2026")
        wb = openpyxl.load_workbook(io.BytesIO(resp.data))
        ws = wb["By Device Type 2026"]
        rows = {r[0]: r[7] for r in ws.iter_rows(min_row=2, values_only=True)}  # column H = 2026-07
        self.assertEqual(rows, {"LCD": 2, "PC": 1})

    def test_device_type_links_to_filtered_manage_assets(self):
        html = self.client.get("/").get_data(as_text=True)
        self.assertIn('href="/assets/?device_name=LCD"', html)
        listing = self.client.get("/assets/?device_name=LCD").get_data(as_text=True)
        self.assertIn("<strong>2</strong> asset(s) match", listing)  # office LCD + CCTV monitor, current month
        self.assertNotIn('data-value="PC"', listing)

    def test_branch_detail_links_to_filtered_lists(self):
        html = self.client.get("/branch/001").get_data(as_text=True)
        # Device name (Month by Month + Device Type Breakdown) and a status count.
        self.assertIn('href="/assets/?branch_no=001&amp;device_name=LCD"', html)
        self.assertIn('href="/assets/?branch_no=001&amp;device_name=LCD&amp;status=USING+LOCAL"', html)
        # CCTV Breakdown -> Manage CCTV.
        self.assertIn('href="/cctv/?branch_no=001&amp;device_name=DVR/CCTV+RECORDER"', html)
        listing = self.client.get("/assets/?branch_no=001&device_name=LCD&status=USING+LOCAL").get_data(as_text=True)
        self.assertIn("<strong>2</strong> asset(s) match", listing)
        cctv = self.client.get("/cctv/?branch_no=001&device_name=DVR/CCTV+RECORDER").get_data(as_text=True)
        self.assertIn("<strong>1</strong> CCTV item(s) match", cctv)

    def test_pages_render(self):
        html = self.client.get("/").get_data(as_text=True)
        self.assertIn("CCTV / DVR (counted separately)", html)
        html = self.client.get("/branch/001").get_data(as_text=True)
        self.assertIn("3 assets currently on record", html)
        self.assertIn("+ 2 CCTV/DVR counted separately", html)
        self.assertEqual(self.client.get("/export").status_code, 200)
        self.assertEqual(self.client.get("/branch/001/export").status_code, 200)


if __name__ == "__main__":
    unittest.main()
